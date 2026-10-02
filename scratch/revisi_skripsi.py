"""Revisi dokumen skripsi (Fase 5):
1. Periode data Juli-November 2026 -> Januari 2026 (fakta data aktual 548 transaksi, 1-31 Jan 2026)
2. Tabel 6 & narasi: Equal Interval -> Tertile (persentil 33,33/66,67, sesuai core_ml.py)
3. PHP + Laravel -> Python + Flask (P348)
4. Tambah narasi: Elbow Method + Silhouette Score via notebook eksperimen (P321)
5. Tambah narasi: artefak model JSON + ambang batas retrain (P348)

Backup sudah dibuat: SafiraAzzahra-2220803049-SKRIPSI.backup-sebelum-revisi.docx
"""
from docx import Document

PATH = r"C:\Users\User\Bismillah Skripsi Safira\SafiraAzzahra-2220803049-SKRIPSI.docx"
doc = Document(PATH)


def set_para_text(p, new_text):
    """Ganti teks paragraf mempertahankan format run pertama."""
    if p.runs:
        p.runs[0].text = new_text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(new_text)


def set_cell_text(cell, new_text):
    p = cell.paragraphs[0]
    set_para_text(p, new_text)
    for extra in cell.paragraphs[1:]:
        extra._element.getparent().remove(extra._element)


def expect(cond, msg):
    if not cond:
        raise AssertionError(msg)
    print(f"  OK: {msg}")


# ---------------------------------------------------------------------------
# 1. Periode data (P176)
# ---------------------------------------------------------------------------
p176 = doc.paragraphs[176]
expect("Juli hingga November tahun 2026" in p176.text, "P176 masih memuat periode lama")
set_para_text(
    p176,
    "Dalam penelitian ini, data yang dimanfaatkan adalah data transaksi pelanggan pada periode "
    "Januari tahun 2026, yang berfungsi sebagai data input dalam perhitungan RFM dan proses clustering."
)
expect("Januari tahun 2026" in p176.text, "P176 periode diubah ke Januari 2026")

# ---------------------------------------------------------------------------
# 2. Narasi Equal Interval -> Tertile (P332)
# ---------------------------------------------------------------------------
p332 = doc.paragraphs[332]
expect("Equal Interval" in p332.text, "P332 masih memuat Equal Interval")
set_para_text(
    p332,
    "Penentuan range kelompok prioritas dilakukan untuk mengelompokkan nilai RFM ke dalam kategori "
    "tinggi, sedang, dan rendah sehingga mempermudah proses interpretasi hasil segmentasi pelanggan. "
    "Pada penelitian ini, range ditentukan menggunakan metode Tertile, yaitu dengan membagi "
    "distribusi nilai RFM pelanggan menjadi tiga kelompok berukuran sama besar berdasarkan nilai "
    "persentil 33,33% dan persentil 66,67% dari data pelanggan yang bertransaksi. Nilai pada "
    "persentil 33,33 ke bawah termasuk kelompok rendah, nilai antara persentil 33,33 sampai "
    "66,67 termasuk kelompok sedang, dan nilai di atas persentil 66,67 termasuk kelompok tinggi. "
    "Untuk Recency berlaku kebalikannya, yaitu nilai recency yang kecil (persentil 33,33 ke bawah) "
    "termasuk kelompok tinggi karena menunjukkan transaksi terbaru. Metode Tertile dipilih karena "
    "distribusi nilai transaksi pelanggan tidak selalu merata, sehingga pembagian berdasarkan "
    "persentil mampu menghasilkan tiga kelompok yang seimbang pada data yang bersifat skew. Hasil "
    "pembagian tersebut digunakan sebagai dasar dalam penentuan kelompok prioritas pelanggan pada sistem."
)
expect("metode Tertile" in p332.text, "P332 diubah ke Tertile")

# ---------------------------------------------------------------------------
# 3. Tabel 6: rentang tetap -> persentil Tertile
# ---------------------------------------------------------------------------
tbl6 = doc.tables[6]
expect(tbl6.rows[0].cells[0].text.strip() == "Kategori", "Tabel[6] adalah Tabel 6 (Kategori/RFM)")

new_cells = {
    (0, 0): "Kategori", (0, 1): "Recency", (0, 2): "Frequency", (0, 3): "Monetary",
    (1, 0): "Tinggi", (1, 1): "\u2264 P33,33", (1, 2): "> P66,67", (1, 3): "> P66,67",
    (2, 0): "Sedang", (2, 1): "P33,33 \u2013 P66,67", (2, 2): "P33,33 \u2013 P66,67", (2, 3): "P33,33 \u2013 P66,67",
    (3, 0): "Rendah", (3, 1): "> P66,67", (3, 2): "\u2264 P33,33", (3, 3): "\u2264 P33,33",
}
for (r, c), teks in new_cells.items():
    set_cell_text(tbl6.rows[r].cells[c], teks)
expect("P33,33" in tbl6.rows[1].cells[1].text, "Tabel 6 baris Tinggi memakai persentil")
expect("0.33" not in tbl6.rows[3].cells[3].text, "Tabel 6 baris Rendah bebas rentang lama")

# ---------------------------------------------------------------------------
# 4+5. P348: Flask + narasi artefak model & retrain
# ---------------------------------------------------------------------------
p348 = doc.paragraphs[348]
expect("PHP dengan framework Laravel" in p348.text, "P348 masih memuat PHP/Laravel")
set_para_text(
    p348,
    "Tahap implementasi dilakukan dengan membangun sistem berbasis web menggunakan bahasa "
    "pemrograman Python dengan framework Flask dan basis data MySQL. Sistem yang telah dibangun "
    "kemudian diuji menggunakan Black-Box Testing dengan teknik Equivalence Partitioning untuk "
    "memastikan seluruh fungsi sistem berjalan dengan baik. Pengujian dilakukan pada fitur "
    "pengelolaan data pelanggan, data transaksi, perhitungan RFM, proses clustering K-Means, serta "
    "penyajian hasil segmentasi pelanggan. Model K-Means yang telah dilatih disimpan dalam format "
    "JSON berisi skala normalisasi dan posisi sentroid, sehingga sistem dapat menilai pelanggan "
    "baru tanpa perlu melatih ulang model. Sistem melakukan pelatihan ulang (retrain) hanya ketika "
    "salah satu ambang batas berikut terpenuhi: penambahan data sebesar 10% dari jumlah data latih, "
    "Silhouette Score di bawah 0,35, lebih dari 10% nilai pelanggan berada di luar rentang skaler "
    "model (drift), atau usia model lebih dari 90 hari."
)
expect("framework Flask" in p348.text, "P348 diubah ke Python/Flask")
expect("retrain" in p348.text, "P348 memuat narasi retrain")

# ---------------------------------------------------------------------------
# 6. P321: narasi Elbow Method + Silhouette Score (notebook eksperimen)
# ---------------------------------------------------------------------------
p321 = doc.paragraphs[321]
expect("evaluasi dan penyajian hasil" in p321.text, "P321 memuat tahapan KDD")
set_para_text(
    p321,
    p321.text.rstrip()
    + " Pada tahap evaluasi pola, kualitas clustering diukur menggunakan Elbow Method untuk "
      "memverifikasi jumlah cluster optimal dan Silhouette Score untuk mengukur keterpisahan "
      "klaster. Evaluasi tersebut dijalankan melalui notebook eksperimen yang mengimpor fungsi "
      "pemrosesan data yang sama dengan sistem, sehingga hasil evaluasi yang diperoleh "
      "konsisten dengan hasil yang ditampilkan oleh sistem."
)
expect("Silhouette Score" in p321.text, "P321 memuat Elbow & Silhouette")

# ---------------------------------------------------------------------------
# Simpan & verifikasi ulang
# ---------------------------------------------------------------------------
doc.save(PATH)
print("\nDokumen disimpan.")

doc2 = Document(PATH)
checks = [
    ("Juli hingga November", False),
    ("Equal Interval", False),
    ("framework Laravel", False),
    ("Januari tahun 2026", True),
    ("metode Tertile", True),
    ("framework Flask", True),
    ("Silhouette Score", True),
    ("retrain", True),
]
print("\n=== Verifikasi setelah simpan ===")
for kata, harapan in checks:
    ada = any(kata in p.text for p in doc2.paragraphs)
    status = "LULUS" if ada == harapan else "GAGAL"
    print(f"  [{status}] '{kata}' ada={ada} (diharapkan {harapan})")
    assert ada == harapan, f"verifikasi gagal: {kata}"
print("Tabel 6:", [[c.text for c in r.cells] for r in doc2.tables[6].rows])
print("\nSEMUA REVISI TERVERIFIKASI")
